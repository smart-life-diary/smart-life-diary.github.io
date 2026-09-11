"""提案書チェック45 のExcel版（自動採点）を書き出す。

出力
  samples/提案書チェック45/提案書チェック45.xlsx

ココナラのコンテンツ出品はHTMLもzipも受け付けない（公式の推奨形式にも無く、
実データ5,012件でも0件）。PDFとXLSXが通る形なので、診断ツールはExcelで作る。

購入者の環境が新しいとは限らないので、FILTER や動的配列は使わない。
0点項目の抽出は「作業列に通し番号を振って INDEX/MATCH で詰める」古い書き方で行う。
Googleスプレッドシートでも LibreOffice でも同じように動く。

色の作法は建設帳票5点セットと揃える。黄色＝入力する欄、グレー＝自動で計算される欄。
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from check45_items import AXES, ITEMS, BANDS, AXIS_ADVICE

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(HERE, "..", "..", "samples", "提案書チェック45"))

INK = "FF182438"
HEAD_BG = "FF1F3B57"
IN_BG = "FFFFF3CC"      # 入力する欄（黄）
CALC_BG = "FFEEF2F6"    # 自動で計算される欄（グレー）
SOFT = "FFE8F1FC"
LINE = "FFD6DCE4"
ACCENT = "FF0070D6"

THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CHOICES = ["できている", "あいまい", "できていない"]
FIRST = 4               # チェックシートの1行目のデータ行


# フォント名は建設帳票5点セットと揃える。あちらは実際に売れて、
# 購入者の環境で問題が出ていない。名前を変える理由が無い。
def f(sz=11, b=False, color=INK):
    return Font(name="Meiryo", size=sz, bold=b, color=color)


def fill(c):
    return PatternFill("solid", fgColor=c)


def put(ws, cell, value, *, font=None, bg=None, align=None, border=True, wrap=False):
    c = ws[cell]
    c.value = value
    c.font = font or f()
    if bg:
        c.fill = fill(bg)
    if border:
        c.border = BORDER
    c.alignment = Alignment(horizontal=align or "left", vertical="center", wrap_text=wrap)
    return c


def sheet_intro(wb):
    ws = wb.create_sheet("はじめに")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 96

    put(ws, "B2", "提案書・営業資料チェック45", font=f(20, True), border=False)
    put(ws, "B3", "出来上がった資料を横に置いて、45項目に答えてください。",
        font=f(11, color="FF6C7C91"), border=False)
    ws.row_dimensions[2].height = 34

    rows = [
        ("使い方", None),
        ("1.  「チェック」シートを開き、D列のプルダウンから3つのうち1つを選びます。", None),
        ("2.  45項目すべて選び終えると、「結果」シートに点数が出ます。", None),
        ("3.  「結果」シートの下に、できていない項目だけが並びます。上から直してください。", None),
        ("", None),
        ("配点", None),
        ("できている 2点 ／ あいまい 1点 ／ できていない 0点。45項目で90点満点です。", None),
        ("構成・情報・表現の3つの軸に15項目ずつ入っていて、軸ごとに30点です。", None),
        ("", None),
        ("色の見分け方", None),
        ("黄色の欄  … あなたが入力する欄です。ここだけ触ります。", IN_BG),
        ("グレーの欄 … 自動で計算される欄です。数式が入っているので触らないでください。", CALC_BG),
        ("", None),
        ("直す順番", None),
        ("いちばん点数の低い軸から手を付けてください。3つを同時に直そうとすると、", None),
        ("表現を整えたあとに構成をやり直すことになり、かえって時間がかかります。", None),
        ("", None),
        ("この45項目について", None),
        ("建設コンサルタントとして15年、公共案件の技術提案書を作成してきた経験から起こしています。", None),
        ("「見やすくしましょう」のような一般論ではなく、通らないときに実際に落ちている箇所を並べました。", None),
    ]
    r = 5
    for text, bg in rows:
        if not text:
            r += 1
            continue
        head = text in ("使い方", "配点", "色の見分け方", "直す順番", "この45項目について")
        c = put(ws, f"B{r}", text, font=f(12, True) if head else f(11),
                bg=bg, border=bool(bg))
        if head:
            c.font = f(12, True, ACCENT)
        r += 1
    return ws


def sheet_check(wb):
    ws = wb.create_sheet("チェック")
    ws.sheet_view.showGridLines = False
    widths = {"A": 5, "B": 8, "C": 62, "D": 15, "E": 7, "F": 9}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v

    put(ws, "A1", "チェック", font=f(16, True), border=False)
    put(ws, "A2", "D列のプルダウンから選んでください。E列とF列は自動です。",
        font=f(10, color="FF6C7C91"), border=False)

    heads = ["No", "軸", "チェック項目", "評価", "点数", "（作業用）"]
    for i, h in enumerate(heads):
        put(ws, f"{get_column_letter(i+1)}3", h, font=f(10, True, "FFFFFFFF"),
            bg=HEAD_BG, align="center")
    ws.row_dimensions[3].height = 22

    dv = DataValidation(type="list", formula1='"{}"'.format(",".join(CHOICES)),
                        allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)

    for i, (ax, text) in enumerate(ITEMS):
        r = FIRST + i
        put(ws, f"A{r}", i + 1, align="center")
        put(ws, f"B{r}", AXES[ax][0], align="center", bg=SOFT)
        put(ws, f"C{r}", text, wrap=True)
        put(ws, f"D{r}", None, bg=IN_BG, align="center")
        dv.add(ws[f"D{r}"])
        # 点数。未回答は空欄のままにして、合計に混ぜない。
        put(ws, f"E{r}",
            f'=IF($D{r}="","",IF($D{r}="{CHOICES[0]}",2,IF($D{r}="{CHOICES[1]}",1,0)))',
            bg=CALC_BG, align="center")
        # できていない項目に上から通し番号を振る。結果シートがこれを拾う。
        put(ws, f"F{r}",
            f'=IF($E{r}=0,COUNTIF($E${FIRST}:$E{r},0),"")',
            bg=CALC_BG, align="center")
        ws.row_dimensions[r].height = 30

    ws.freeze_panes = "A4"
    return ws


def sheet_result(wb):
    ws = wb.create_sheet("結果")
    ws.sheet_view.showGridLines = False
    for k, v in {"A": 3, "B": 22, "C": 14, "D": 12, "E": 60}.items():
        ws.column_dimensions[k].width = v

    last = FIRST + len(ITEMS) - 1
    put(ws, "B2", "結果", font=f(20, True), border=False)
    ws.row_dimensions[2].height = 34

    put(ws, "B4", "総合", font=f(12, True), bg=SOFT, align="center")
    put(ws, "C4", f"=SUM(チェック!$E${FIRST}:$E${last})", font=f(20, True, ACCENT),
        bg=CALC_BG, align="center")
    put(ws, "D4", "／ 90", font=f(11, color="FF6C7C91"), bg=CALC_BG, align="center")
    cond = "".join(f'IF($C$4>={lo},"{msg}",' for lo, msg, _ in BANDS[:-1])
    put(ws, "E4", f'={cond}"{BANDS[-1][1]}"' + ")" * (len(BANDS) - 1),
        font=f(12, True), bg=CALC_BG)
    ws.row_dimensions[4].height = 30

    cond2 = "".join(f'IF($C$4>={lo},"{d}",' for lo, _, d in BANDS[:-1])
    put(ws, "E5", f'={cond2}"{BANDS[-1][2]}"' + ")" * (len(BANDS) - 1),
        font=f(10, color="FF6C7C91"), bg=CALC_BG, wrap=True)
    ws.row_dimensions[5].height = 30

    put(ws, "B7", "軸別", font=f(12, True, ACCENT), border=False)
    for i, (name, note, _c, _d) in enumerate(AXES):
        r = 8 + i
        put(ws, f"B{r}", name, font=f(11, True), bg=SOFT, align="center")
        put(ws, f"C{r}",
            f'=SUMIF(チェック!$B${FIRST}:$B${last},"{name}",チェック!$E${FIRST}:$E${last})',
            font=f(13, True), bg=CALC_BG, align="center")
        put(ws, f"D{r}", "／ 30", font=f(10, color="FF6C7C91"), bg=CALC_BG, align="center")
        put(ws, f"E{r}", note, font=f(10, color="FF6C7C91"), bg=CALC_BG)
        ws.row_dimensions[r].height = 24

    put(ws, "B12", "いちばん低い軸", font=f(12, True, ACCENT), border=False)
    low = 'INDEX({' + ",".join(f'"{a[0]}"' for a in AXES) + '},MATCH(MIN($C$8:$C$10),$C$8:$C$10,0))'
    put(ws, "C12", f"={low}", font=f(12, True), bg=CALC_BG, align="center")
    put(ws, "E12", "ここから手を付けてください。", font=f(10, color="FF6C7C91"), bg=CALC_BG)

    r = 13
    for i in range(4):
        cond3 = "".join(
            f'IF($C$12="{AXES[k][0]}","{AXIS_ADVICE[k][i]}",' for k in range(len(AXES) - 1))
        put(ws, f"E{r}",
            f'={cond3}"{AXIS_ADVICE[-1][i]}"' + ")" * (len(AXES) - 1),
            font=f(10), bg=CALC_BG, wrap=True)
        put(ws, f"D{r}", "→", font=f(10, True, ACCENT), bg=CALC_BG, align="center")
        ws.row_dimensions[r].height = 28
        r += 1

    r += 1
    put(ws, f"B{r}", "できていない項目", font=f(12, True, ACCENT), border=False)
    put(ws, f"C{r}", f'=COUNTIF(チェック!$E${FIRST}:$E${last},0)&" 件"',
        font=f(11, True), border=False)
    r += 1
    put(ws, f"B{r}", "上から順に潰してください。", font=f(10, color="FF6C7C91"), border=False)
    r += 1
    top = r
    for i in range(len(ITEMS)):
        rr = top + i
        put(ws, f"B{rr}",
            f'=IFERROR(INDEX(チェック!$A${FIRST}:$A${last},'
            f'MATCH({i+1},チェック!$F${FIRST}:$F${last},0)),"")',
            bg=CALC_BG, align="center")
        put(ws, f"C{rr}",
            f'=IFERROR(INDEX(チェック!$B${FIRST}:$B${last},'
            f'MATCH({i+1},チェック!$F${FIRST}:$F${last},0)),"")',
            bg=CALC_BG, align="center")
        ws.merge_cells(f"D{rr}:E{rr}")
        put(ws, f"D{rr}",
            f'=IFERROR(INDEX(チェック!$C${FIRST}:$C${last},'
            f'MATCH({i+1},チェック!$F${FIRST}:$F${last},0)),"")',
            bg=CALC_BG)
        ws[f"E{rr}"].border = BORDER
        ws.row_dimensions[rr].height = 24
    return ws


def build():
    wb = Workbook()
    wb.remove(wb.active)
    sheet_intro(wb)
    sheet_check(wb)
    sheet_result(wb)
    # 生成直後の .xlsx には計算結果が入らない。開いた時点で必ず再計算させる。
    wb.calculation.fullCalcOnLoad = True
    return wb


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    p = os.path.join(OUT, "提案書チェック45.xlsx")
    build().save(p)
    print(p, os.path.getsize(p) // 1024, "KB")

    import zipfile
    with zipfile.ZipFile(p) as z:
        x = z.read("xl/workbook.xml").decode()
    print("fullCalcOnLoad:", 'fullCalcOnLoad="1"' in x)
